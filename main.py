import io
import os
import re
import json
from copy import deepcopy
from typing import Any, Dict, Optional, Iterable, Tuple, List

from fastapi import FastAPI, UploadFile, File, Form, Header, HTTPException
from fastapi.responses import StreamingResponse

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

app = FastAPI(title="docx-format-api", version="0.2.0")

PLACEHOLDER = "{{CONTENT}}"
API_KEY = os.getenv("API_KEY", "").strip() or None

# 默认：删除占位符之后的模板内容（避免模板示例页残留）
DELETE_TAIL_AFTER_PLACEHOLDER = os.getenv("DELETE_TAIL_AFTER_PLACEHOLDER", "true").lower() in ("1", "true", "yes")


# -----------------------
# Utils: parse rules (方案A)
# -----------------------
def parse_rules(raw: Optional[str]) -> Optional[Dict[str, Any]]:
    if not raw:
        return None
    raw = raw.strip()

    # 兼容：有人把整个 LLM 输出对象都传来了 {"text":"..."}
    try:
        maybe = json.loads(raw)
        if isinstance(maybe, dict) and isinstance(maybe.get("text"), str):
            raw = maybe["text"].strip()
    except Exception:
        pass

    # 第一次 loads：可能得到 dict，也可能得到 str(里面还是JSON)
    try:
        obj = json.loads(raw)
    except Exception:
        # 兜底截取
        l = raw.find("{")
        r = raw.rfind("}")
        if l != -1 and r != -1 and r > l:
            obj = json.loads(raw[l:r + 1])
        else:
            raise

    if isinstance(obj, str):
        obj = json.loads(obj)

    return obj if isinstance(obj, dict) else None


# -----------------------
# Utils: iterate paragraphs (including tables)
# -----------------------
def iter_paragraphs_in_cell(cell) -> Iterable[Paragraph]:
    for p in cell.paragraphs:
        yield p
    for t in cell.tables:
        for row in t.rows:
            for c in row.cells:
                yield from iter_paragraphs_in_cell(c)

def iter_all_paragraphs(doc: Document) -> Iterable[Paragraph]:
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from iter_paragraphs_in_cell(cell)


# -----------------------
# Utils: style helpers (set pStyle by style_id)
# -----------------------
def style_name_to_id_map(doc: Document) -> Dict[str, str]:
    out = {}
    for s in doc.styles:
        try:
            if getattr(s, "name", None) and getattr(s, "style_id", None):
                out[s.name] = s.style_id
        except Exception:
            continue
    return out

def pick_style_id(template_doc: Document, candidates: List[str]) -> Optional[str]:
    m = style_name_to_id_map(template_doc)
    for name in candidates:
        if name in m:
            return m[name]
    return None

def set_paragraph_style_id(p: Paragraph, style_id: str) -> None:
    if not style_id:
        return
    p_elm = p._p
    pPr = p_elm.get_or_add_pPr()
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        pStyle = OxmlElement("w:pStyle")
        pPr.insert(0, pStyle)
    pStyle.set(qn("w:val"), style_id)


# -----------------------
# Rules compilation
# -----------------------
def compile_detect(detect: str) -> Tuple[str, Any]:
    detect = (detect or "").strip()
    if detect.lower().startswith("regex:"):
        pat = detect[6:].strip()
        return ("regex", re.compile(pat))
    if detect.lower().startswith("contains:"):
        keys = [k.strip() for k in detect[9:].split("|") if k.strip()]
        return ("contains", keys)
    return ("contains", [detect])

def match_detect(compiled: Tuple[str, Any], text: str) -> bool:
    kind, obj = compiled
    if kind == "regex":
        return obj.search(text) is not None
    return any(k in text for k in obj)

def normalize_rules(rules_obj: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    # 只取我们用得上的字段
    out = {"formatting_intent": {"heading_rules": [], "special_rules": [], "body_style_name": None}}
    if not rules_obj:
        return out
    fi = rules_obj.get("formatting_intent") or {}
    out["formatting_intent"]["body_style_name"] = fi.get("body_style_name")
    out["formatting_intent"]["heading_rules"] = fi.get("heading_rules") or []
    out["formatting_intent"]["special_rules"] = fi.get("special_rules") or []
    return out


# -----------------------
# Apply styles to source (content unchanged)
# -----------------------
RX_H1 = re.compile(r"^(第[一二三四五六七八九十]+章)\b|^\d+\s*[\.、]\s+\S")
RX_H2 = re.compile(r"^\d+\.\d+\s+\S")
RX_H3 = re.compile(r"^\d+\.\d+\.\d+\s+\S")
RX_REF_TITLE = re.compile(r"^(参考文献|REFERENCES)\b", re.IGNORECASE)
RX_FIG_CAP = re.compile(r"^图\s*\d+|^Figure\s*\d+", re.IGNORECASE)
RX_TAB_CAP = re.compile(r"^表\s*\d+|^Table\s*\d+", re.IGNORECASE)

def get_template_style_ids(template_doc: Document) -> Dict[str, Optional[str]]:
    # 优先用你模板里的 CM*，否则回退 Word 默认
    return {
        "body": pick_style_id(template_doc, ["CMcontent", "正文", "Normal", "Body Text"]),
        "h1": pick_style_id(template_doc, ["CMheading1", "标题1", "标题 1", "Heading 1"]),
        "h2": pick_style_id(template_doc, ["CMheading2", "标题2", "标题 2", "Heading 2"]),
        "h3": pick_style_id(template_doc, ["CMheading3", "标题3", "标题 3", "Heading 3"]),
        "fig_caption": pick_style_id(template_doc, ["CMcaption", "题注", "Caption"]),
        "tab_caption": pick_style_id(template_doc, ["CMcaptionTable", "题注", "Caption"]),
        "ref_item": pick_style_id(template_doc, ["CMreflist", "参考文献", "Bibliography", "References"]),
    }

def apply_styles_in_place(template_doc: Document, source_doc: Document, rules_obj: Optional[Dict[str, Any]]) -> None:
    style_ids = get_template_style_ids(template_doc)

    rules = normalize_rules(rules_obj)
    fi = rules["formatting_intent"]
    heading_rules = fi.get("heading_rules") or []
    special_rules = fi.get("special_rules") or []

    compiled_heading = []
    for r in heading_rules:
        try:
            lvl = int(r.get("level"))
            compiled_heading.append((lvl, compile_detect(r.get("detect") or ""), r.get("style_name")))
        except Exception:
            pass
    compiled_heading.sort(key=lambda x: x[0], reverse=True)  # 先匹配 3 再 2 再 1

    compiled_special = []
    for r in special_rules:
        try:
            compiled_special.append((r.get("type"), compile_detect(r.get("detect") or ""), r.get("style_name")))
        except Exception:
            pass

    # special_rules 里我们只用这些 type（其余忽略）
    type_to_key = {
        "caption_figure": "fig_caption",
        "caption_table": "tab_caption",
    }

    in_references = False
    for p in iter_all_paragraphs(source_doc):
        text = (p.text or "").strip()
        if not text:
            continue

        # 参考文献段开始：之后都套 ref_item
        if RX_REF_TITLE.search(text):
            in_references = True
            if style_ids.get("h1"):
                set_paragraph_style_id(p, style_ids["h1"])
            continue

        if in_references and style_ids.get("ref_item"):
            set_paragraph_style_id(p, style_ids["ref_item"])
            continue

        # special_rules（只先处理图/表题注）
        applied = False
        for typ, det, style_name in compiled_special:
            if typ in type_to_key and match_detect(det, text):
                key = type_to_key[typ]
                sid = style_ids.get(key)
                if sid:
                    set_paragraph_style_id(p, sid)
                    applied = True
                break
        if applied:
            continue

        # 题注启发式兜底
        if style_ids.get("fig_caption") and RX_FIG_CAP.search(text):
            set_paragraph_style_id(p, style_ids["fig_caption"])
            continue
        if style_ids.get("tab_caption") and RX_TAB_CAP.search(text):
            set_paragraph_style_id(p, style_ids["tab_caption"])
            continue

        # heading_rules（如果给了）
        if compiled_heading:
            matched = False
            for lvl, det, style_name in compiled_heading:
                if match_detect(det, text):
                    # 如果 style_name 在模板里存在，就优先用它
                    if style_name and style_name in style_name_to_id_map(template_doc):
                        set_paragraph_style_id(p, style_name_to_id_map(template_doc)[style_name])
                    else:
                        sid = style_ids.get(f"h{lvl}")
                        if sid:
                            set_paragraph_style_id(p, sid)
                    matched = True
                    break
            if matched:
                continue

        # MVP 启发式
        if style_ids.get("h3") and RX_H3.search(text):
            set_paragraph_style_id(p, style_ids["h3"])
        elif style_ids.get("h2") and RX_H2.search(text):
            set_paragraph_style_id(p, style_ids["h2"])
        elif style_ids.get("h1") and RX_H1.search(text):
            set_paragraph_style_id(p, style_ids["h1"])
        else:
            if style_ids.get("body"):
                set_paragraph_style_id(p, style_ids["body"])


# -----------------------
# Insert into template at placeholder
# -----------------------
def find_placeholder_paragraph(doc: Document) -> Optional[Paragraph]:
    for p in doc.paragraphs:
        if (p.text or "").strip() == PLACEHOLDER:
            return p
    return None

def delete_body_elements_after(anchor_p: Paragraph, doc: Document) -> None:
    """Delete all body children after anchor paragraph (in the main body only)."""
    body = doc.element.body
    anchor_elm = anchor_p._p
    children = list(body.iterchildren())
    # Find index of anchor element in body children
    try:
        idx = children.index(anchor_elm)
    except ValueError:
        return
    # Remove after idx, but keep sectPr
    for child in children[idx + 1:]:
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)

def insert_source_after(anchor_p: Paragraph, source_doc: Document) -> None:
    anchor_elm = anchor_p._p
    parent = anchor_elm.getparent()
    idx = parent.index(anchor_elm)

    for child in source_doc.element.body.iterchildren():
        if child.tag.endswith("}sectPr"):
            continue
        parent.insert(idx + 1, deepcopy(child))
        idx += 1


@app.get("/health")
def health():
    return {"ok": True}


@app.post("/format")
async def format_docx(
    template: UploadFile = File(...),
    source: UploadFile = File(...),
    rules: Optional[str] = Form(None),
    x_api_key: Optional[str] = Header(None, convert_underscores=False),
):
    if API_KEY and (not x_api_key or x_api_key.strip() != API_KEY):
        raise HTTPException(status_code=401, detail="Unauthorized")

    t_bytes = await template.read()
    s_bytes = await source.read()

    try:
        tdoc = Document(io.BytesIO(t_bytes))
        sdoc = Document(io.BytesIO(s_bytes))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to open docx: {e}")

    rules_obj = None
    if rules:
        try:
            rules_obj = parse_rules(rules)
        except Exception:
            rules_obj = None

    # 1) 先在 source 上套模板样式（只改样式，不改文字）
    apply_styles_in_place(tdoc, sdoc, rules_obj)

    # 2) 在模板中找到插入点
    anchor = find_placeholder_paragraph(tdoc)
    if anchor is None:
        # 找不到占位符，就插到末尾（但会不理想）
        anchor = tdoc.add_paragraph("")
    else:
        # 删除占位符后面的模板内容（防止示例页残留）
        if DELETE_TAIL_AFTER_PLACEHOLDER:
            delete_body_elements_after(anchor, tdoc)
        # 清空占位符那一行文字
        anchor.text = ""

    # 3) 插入 source 的块级内容（保留表格等）
    insert_source_after(anchor, sdoc)

    out = io.BytesIO()
    tdoc.save(out)
    out.seek(0)

    headers = {"Content-Disposition": 'attachment; filename="result.docx"'}
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=headers,
    )
